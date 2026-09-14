"""
src/sovereign/vision/register.py
Structured Engineering Change Register Generator.
Layer 2: Vision-Language Model semantic analysis via BaseModelAdapter.vision_infer(),
Markdown tabular formatting, and styled Word (.docx) table deliverable export.
Complies with MMOD-03, DELV-01, T-04-04 (Path Fencing), and ASVS V5.1/V10.3.
"""
from datetime import datetime, timezone
import io
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Union
import uuid

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, Field
from PIL import Image

from sovereign.deliverables.docx import (
    COLOR_BODY_CHARCOAL,
    COLOR_PRIMARY_NAVY,
    COLOR_SECONDARY_SLATE,
    COLOR_WHITE,
    HEX_PRIMARY_NAVY,
    HEX_WHITE,
    HEX_ZEBRA_FILL,
    set_cell_margins,
    set_cell_shading,
    validate_output_path,
)
from sovereign.lineage.engine import LineageEngine
from sovereign.lineage.models import DeliverableInfo, LineageMetadata, ModelInfo
from sovereign.models.adapter import BaseModelAdapter
from sovereign.vision.diff import DiffBoundingBox, DiffChangeType, DrawingDiffEngine

# Industrial color styling for status/risk badges
COLOR_ADDED_HEX = "00C853"
COLOR_REMOVED_HEX = "D50000"
COLOR_MODIFIED_HEX = "FFD600"
COLOR_RISK_HIGH_HEX = "FEE2E2"    # Light red fill
COLOR_RISK_MED_HEX = "FEF3C7"     # Light amber fill
COLOR_RISK_LOW_HEX = "DCFCE7"     # Light green fill


class ChangeRegisterItem(BaseModel):
    """Structured line item detailing a verified revision change to an engineering drawing."""
    item_id: int = Field(..., description="Matches box_id on visual overlay and numbered badge")
    change_type: DiffChangeType = Field(..., description="ADDED, REMOVED, or MODIFIED")
    equipment_tag: str = Field(default="UNKNOWN", description="Identified equipment or instrumentation tag")
    description: str = Field(..., description="Engineering description of physical drawing modification")
    rev_a_state: str = Field(default="Not present", description="Baseline state in Revision A")
    rev_b_state: str = Field(default="Active", description="Updated state in Revision B")
    safety_risk: str = Field(default="LOW", description="Safety risk classification: LOW, MEDIUM, or HIGH")
    box_2d: List[float] = Field(..., description="Normalized spatial coordinates [ymin, xmin, ymax, xmax]")


class ChangeRegister(BaseModel):
    """Complete structured revision change register for an engineering drawing pair."""
    drawing_title: str = Field(default="P&ID Revision Comparison")
    rev_a_name: str = Field(default="Revision A")
    rev_b_name: str = Field(default="Revision B")
    analysis_iso: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    items: List[ChangeRegisterItem] = Field(default_factory=list)


class ChangeRegisterGenerator:
    """
    Generates structured Change Register documentation by combining Layer 1 spatial diff boxes
    with Layer 2 Vision-Language reasoning via BaseModelAdapter.vision_infer().
    """

    @classmethod
    def _extract_semantic_details(
        cls,
        box: DiffBoundingBox,
        vlm_response_text: str,
        vision_regions: Optional[List[Any]] = None,
    ) -> Dict[str, str]:
        """
        Parses VLM inference text or structured vision regions into structured change attributes.
        """
        tag = None
        desc = None
        safety_risk = "LOW"
        rev_a_state = "Not present" if box.change_type == DiffChangeType.ADDED else "Baseline configuration"
        rev_b_state = "Removed from service" if box.change_type == DiffChangeType.REMOVED else "Active in service"

        # 1. Check vision_regions if provided by VLM adapter
        if vision_regions:
            for vr in vision_regions:
                label = getattr(vr, "label", "")
                if label:
                    tag = label.replace("_", "-").upper()
                    break

        # 2. Extract equipment tags matching industrial regex patterns (e.g. CV-103, P-101, V-102, PT-104)
        if not tag:
            match = re.search(r"\b([A-Z]{1,4}-\d{2,4}[A-Z]?)\b", vlm_response_text)
            if match:
                tag = match.group(1)

        # 3. Assess safety risk keywords in VLM output
        text_lower = vlm_response_text.lower()
        if any(k in text_lower for k in ["high risk", "hazardous", "relief", "overpressure", "critical"]):
            safety_risk = "HIGH"
        elif any(k in text_lower for k in ["medium risk", "isolation", "bypass", "control valve", "re-route"]):
            safety_risk = "MEDIUM"

        # 4. Formulate engineering description
        if not desc:
            if box.change_type == DiffChangeType.ADDED:
                desc = f"Introduced component {tag or f'TAG-{box.box_id:02d}'} into process stream"
            elif box.change_type == DiffChangeType.REMOVED:
                desc = f"Decommissioned/removed component {tag or f'TAG-{box.box_id:02d}'} from schematic"
            else:
                desc = f"Modified geometry or specifications for {tag or f'TAG-{box.box_id:02d}'}"

        if not tag:
            tag = f"TAG-{box.box_id:02d}"

        return {
            "equipment_tag": tag,
            "description": desc,
            "rev_a_state": rev_a_state,
            "rev_b_state": rev_b_state,
            "safety_risk": safety_risk,
        }

    @classmethod
    def generate_register(
        cls,
        diff_boxes: List[DiffBoundingBox],
        model_adapter: Optional[BaseModelAdapter],
        img_a_path: Union[str, Path],
        img_b_path: Union[str, Path],
        drawing_title: str = "P&ID Revision Comparison",
    ) -> ChangeRegister:
        """
        Generates a complete ChangeRegister by querying local VLM reasoning over each diff region.

        Args:
            diff_boxes: Layer 1 spatial difference bounding boxes.
            model_adapter: Model adapter providing vision_infer capability (e.g. MockModelAdapter, SGLang/vLLM).
            img_a_path: Baseline drawing Rev A path.
            img_b_path: Updated drawing Rev B path.
            drawing_title: Name/title of the drawing.

        Returns:
            Populated ChangeRegister Pydantic model.
        """
        path_a = Path(img_a_path).resolve()
        path_b = Path(img_b_path).resolve()

        items: List[ChangeRegisterItem] = []

        for box in diff_boxes:
            # Crop the diff region from Rev B (or Rev A if removed)
            target_path = path_b if box.change_type != DiffChangeType.REMOVED else path_a
            crop_img = DrawingDiffEngine.crop_diff_region(target_path, box.box_2d, padding_px=20)

            # Convert crop to PNG bytes
            buf = io.BytesIO()
            crop_img.save(buf, format="PNG")
            crop_bytes = buf.getvalue()

            vlm_text = ""
            vision_regions = None

            if model_adapter is not None:
                prompt = (
                    f"Analyze this engineering drawing crop for revision change #{box.box_id} ({box.change_type.value}). "
                    f"Extract the equipment tag (e.g. CV-101, P-102), describe the modification, and evaluate safety implications."
                )
                try:
                    response = model_adapter.vision_infer(crop_bytes, prompt)
                    vlm_text = response.text or ""
                    vision_regions = response.vision_regions
                except Exception as exc:
                    vlm_text = f"VLM reasoning fallback: {exc}"

            details = cls._extract_semantic_details(box, vlm_text, vision_regions)

            items.append(
                ChangeRegisterItem(
                    item_id=box.box_id,
                    change_type=box.change_type,
                    equipment_tag=details["equipment_tag"],
                    description=details["description"],
                    rev_a_state=details["rev_a_state"],
                    rev_b_state=details["rev_b_state"],
                    safety_risk=details["safety_risk"],
                    box_2d=box.box_2d,
                )
            )

        return ChangeRegister(
            drawing_title=drawing_title,
            rev_a_name=path_a.name,
            rev_b_name=path_b.name,
            items=items,
        )

    @classmethod
    def export_markdown(cls, register: ChangeRegister) -> str:
        """
        Formats the ChangeRegister as a clean Markdown summary table.
        """
        lines = [
            f"# Engineering Change Register: {register.drawing_title}",
            "",
            f"- **Baseline Drawing (Rev A):** `{register.rev_a_name}`",
            f"- **Updated Drawing (Rev B):** `{register.rev_b_name}`",
            f"- **Analysis Timestamp:** `{register.analysis_iso}`",
            f"- **Total Identified Changes:** `{len(register.items)}`",
            "",
            "| Item # | Change Type | Equipment Tag | Description | Rev A State | Rev B State | Safety Risk |",
            "|:---:|:---:|:---:|:---|:---|:---|:---:|",
        ]

        for item in register.items:
            lines.append(
                f"| {item.item_id} | {item.change_type.value} | `{item.equipment_tag}` | {item.description} | {item.rev_a_state} | {item.rev_b_state} | {item.safety_risk} |"
            )

        lines.append("")
        return "\n".join(lines)

    @classmethod
    def export_word_table(
        cls,
        register: ChangeRegister,
        docx_path: Union[str, Path],
        allowed_root: Optional[Union[str, Path]] = None,
        lineage_metadata: Optional[LineageMetadata] = None,
    ) -> Path:
        """
        Appends or creates an engineering Change Register table in a Microsoft Word (.docx) document.
        Employs corporate typography, zebra styling, risk highlighting, and LineageEngine packaging.

        Args:
            register: The populated ChangeRegister model.
            docx_path: Destination path for .docx file.
            allowed_root: Security boundary for path fencing.
            lineage_metadata: Optional lineage metadata.

        Returns:
            Resolved Path to the saved .docx file.
        """
        resolved_path = validate_output_path(docx_path, allowed_root)
        resolved_path.parent.mkdir(parents=True, exist_ok=True)

        if resolved_path.exists():
            doc = docx.Document(str(resolved_path))
        else:
            doc = docx.Document()
            for section in doc.sections:
                section.top_margin = Inches(1.0)
                section.bottom_margin = Inches(1.0)
                section.left_margin = Inches(1.0)
                section.right_margin = Inches(1.0)

            # Document Title
            p_title = doc.add_paragraph()
            r_title = p_title.add_run("P&ID Revision Verification & Change Register")
            r_title.font.name = "Calibri"
            r_title.font.size = Pt(18)
            r_title.font.bold = True
            r_title.font.color.rgb = COLOR_PRIMARY_NAVY
            p_title.paragraph_format.space_after = Pt(4)

        # Section Heading
        h_sec = doc.add_paragraph()
        r_sec = h_sec.add_run(f"Engineering Change Register: {register.drawing_title}")
        r_sec.font.name = "Calibri"
        r_sec.font.size = Pt(14)
        r_sec.font.bold = True
        r_sec.font.color.rgb = COLOR_PRIMARY_NAVY
        h_sec.paragraph_format.space_after = Pt(4)

        # Meta paragraph
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(8)
        r_meta = p_meta.add_run(
            f"Comparison between Rev A ({register.rev_a_name}) and Rev B ({register.rev_b_name})  |  "
            f"Total Changes: {len(register.items)}  |  "
            f"Generated: {register.analysis_iso[:19].replace('T', ' ')} UTC"
        )
        r_meta.font.name = "Calibri"
        r_meta.font.size = Pt(9.5)
        r_meta.font.color.rgb = COLOR_SECONDARY_SLATE

        # Build Table
        # Columns: Item #, Type, Tag, Description, Rev A State, Rev B State, Safety Risk
        headers = [
            ("Item #", Inches(0.7)),
            ("Type", Inches(1.0)),
            ("Tag", Inches(1.1)),
            ("Description", Inches(2.3)),
            ("Rev A State", Inches(1.1)),
            ("Rev B State", Inches(1.1)),
            ("Safety Risk", Inches(1.0)),
        ]

        table = doc.add_table(rows=1 + len(register.items), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        # Format Header Row
        hdr_row = table.rows[0]
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))

        for idx, (hdr_text, col_width) in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = col_width
            set_cell_shading(cell, HEX_PRIMARY_NAVY)
            set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 1, 6] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(hdr_text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_WHITE

        # Format Data Rows
        for row_idx, item in enumerate(register.items, start=1):
            row = table.rows[row_idx]
            is_odd = (row_idx % 2 == 1)
            row_fill = HEX_ZEBRA_FILL if is_odd else HEX_WHITE

            values = [
                str(item.item_id),
                item.change_type.value,
                item.equipment_tag,
                item.description,
                item.rev_a_state,
                item.rev_b_state,
                item.safety_risk,
            ]

            for c_idx, val in enumerate(values):
                cell = row.cells[c_idx]
                cell.width = headers[c_idx][1]
                set_cell_shading(cell, row_fill)
                set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 6] else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9.0)

                # Special coloring for Item # and Type
                if c_idx == 0:
                    run.font.bold = True
                elif c_idx == 1:
                    run.font.bold = True
                    if val == "ADDED":
                        run.font.color.rgb = RGBColor(0, 150, 60)
                    elif val == "REMOVED":
                        run.font.color.rgb = RGBColor(180, 0, 0)
                    else:
                        run.font.color.rgb = RGBColor(180, 130, 0)
                elif c_idx == 6:
                    run.font.bold = True
                    if val == "HIGH":
                        set_cell_shading(cell, COLOR_RISK_HIGH_HEX)
                        run.font.color.rgb = RGBColor(180, 0, 0)
                    elif val == "MEDIUM":
                        set_cell_shading(cell, COLOR_RISK_MED_HEX)
                        run.font.color.rgb = RGBColor(180, 110, 0)
                    else:
                        set_cell_shading(cell, COLOR_RISK_LOW_HEX)
                        run.font.color.rgb = RGBColor(0, 120, 40)
                else:
                    run.font.color.rgb = COLOR_BODY_CHARCOAL

        # Save Word document
        doc.save(resolved_path)

        # Lineage packaging
        meta = lineage_metadata
        if meta is None:
            meta = LineageMetadata(
                deliverable=DeliverableInfo(
                    file_name=resolved_path.name,
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
                session_id=f"session-{uuid.uuid4().hex[:8]}",
                user_id="engineer",
                user_role="Lead Engineer",
                model=ModelInfo(model_id="sovereign-vision-register-v1", weights_sha256="0" * 64),
                audit_start_seq=0,
                audit_end_seq=1,
                audit_tip_hash="0" * 64,
                custom_attributes={
                    "drawing_title": register.drawing_title,
                    "changes_count": len(register.items),
                },
            )

        LineageEngine.package_deliverable(resolved_path, meta)

        return resolved_path
