"""
tests/test_demos/test_demo_c.py
End-to-End Automated Integration Test Suite for Demo Workflow C (P&ID Revision Comparison / DEMO-03).
Tests dual-revision P&ID comparison, spatial differencing detecting exactly 3 regions
(ADDED, REMOVED, MODIFIED), color-coded visual overlay generation with embedded W3C PROV-O lineage,
and structured Change Register export in Word (.docx) and sanitized CSV (Threat T-06-04).
"""
import csv
from pathlib import Path
import docx
import pytest
from PIL import Image

from sovereign.lineage.injectors.png import extract_png_lineage
from sovereign.vision.diff import DiffBoundingBox, DiffChangeType, DrawingDiffEngine
from sovereign.vision.overlay import VisualOverlayRenderer
from sovereign.vision.register import (
    ChangeRegister,
    ChangeRegisterGenerator,
    ChangeRegisterItem,
)


def sanitize_csv_cell(val: object) -> str:
    """
    Sanitizes tabular text fields against spreadsheet formula injection (Threat T-06-04).
    Prefixes a single quote (') if the text begins with '=', '+', '-', or '@'.
    """
    text = str(val) if val is not None else ""
    if text and text[0] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def export_change_register_csv(register: ChangeRegister, csv_path: Path) -> Path:
    """
    Exports a ChangeRegister model to a structured, formula-sanitized CSV file.
    """
    csv_path = Path(csv_path).resolve()
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    headers = [
        "Item #",
        "Change Type",
        "Equipment Tag",
        "Description",
        "Rev A State",
        "Rev B State",
        "Safety Risk",
    ]

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for item in register.items:
            writer.writerow([
                sanitize_csv_cell(item.item_id),
                sanitize_csv_cell(item.change_type.value),
                sanitize_csv_cell(item.equipment_tag),
                sanitize_csv_cell(item.description),
                sanitize_csv_cell(item.rev_a_state),
                sanitize_csv_cell(item.rev_b_state),
                sanitize_csv_cell(item.safety_risk),
            ])

    return csv_path


def test_demo_c_diff_detection(samples_dir: Path):
    """
    Validates DrawingDiffEngine revision comparison on samples/pid_drawing_rev_a.png and rev_b.png:
    - Detects exactly 3 discrete diff regions.
    - Change types include ADDED, REMOVED, and MODIFIED.
    - Normalized spatial coordinates box_2d in [0.0, 1.0].
    """
    path_a = samples_dir / "pid_drawing_rev_a.png"
    path_b = samples_dir / "pid_drawing_rev_b.png"

    assert path_a.exists() and path_b.exists()

    boxes = DrawingDiffEngine.compare_revisions(path_a, path_b)

    assert len(boxes) == 3, f"Expected exactly 3 diff regions, got {len(boxes)}"

    detected_types = {b.change_type for b in boxes}
    assert DiffChangeType.ADDED in detected_types
    assert DiffChangeType.REMOVED in detected_types
    assert DiffChangeType.MODIFIED in detected_types

    for box in boxes:
        ymin, xmin, ymax, xmax = box.box_2d
        assert 0.0 <= ymin < ymax <= 1.0
        assert 0.0 <= xmin < xmax <= 1.0
        assert box.pixel_count >= 20
        assert 0.0 <= box.confidence <= 1.0


def test_demo_c_visual_overlay_rendering(samples_dir: Path, tmp_path: Path):
    """
    Validates VisualOverlayRenderer generating high-contrast diff overlay:
    - Renders pid_diff_overlay.png in output directory.
    - Dimensions match 1600 x 1200 RGB.
    - Contains title block, legend swatches, numbered circular badges.
    - Embeds W3C PROV-O lineage into native PNG tEXt chunk ('sovereign_lineage').
    - Sibling sidecar .lineage.json is generated.
    """
    path_a = samples_dir / "pid_drawing_rev_a.png"
    path_b = samples_dir / "pid_drawing_rev_b.png"
    out_dir = tmp_path / "output"
    out_png = out_dir / "pid_diff_overlay.png"

    boxes = DrawingDiffEngine.compare_revisions(path_a, path_b)
    assert len(boxes) == 3

    result_path = VisualOverlayRenderer.render_diff_overlay(
        img_b_path=path_b,
        diff_boxes=boxes,
        output_path=out_png,
        drawing_title="P&ID Unit 04 Hydrogen Header & Vessel V-301",
        allowed_root=tmp_path,
    )

    assert result_path.exists()
    assert result_path == out_png.resolve()
    assert result_path.stat().st_size > 10_000  # Multi-kilobyte rendered deliverable

    # Verify PNG image properties
    with Image.open(result_path) as img:
        assert img.size == (1600, 1200)
        assert img.mode == "RGB"

    # Verify embedded W3C PROV-O lineage in PNG tEXt chunk
    image_bytes = result_path.read_bytes()
    lineage_str = extract_png_lineage(image_bytes)
    assert lineage_str is not None, "Embedded 'sovereign_lineage' tEXt chunk missing from PNG"
    assert "sovereign-vision-diff-v1" in lineage_str
    assert "pid_diff_overlay.png" in lineage_str

    # Sibling sidecar manifest
    sidecar_path = Path(f"{result_path}.lineage.json")
    assert sidecar_path.exists()
    sidecar_content = sidecar_path.read_text(encoding="utf-8")
    assert "sovereign-vision-diff-v1" in sidecar_content


def test_demo_c_change_register_export(samples_dir: Path, tmp_path: Path):
    """
    Validates ChangeRegister compilation and export:
    - Assigns risk classifications: HIGH for RV-402, MEDIUM for PT-201, LOW for V-102.
    - Exports structured Word change register (pid_revision_change_register.docx).
    - Exports structured CSV change register (pid_revision_change_register.csv).
    - Validates spreadsheet formula injection sanitization (=, +, -, @) per Threat T-06-04.
    """
    path_a = samples_dir / "pid_drawing_rev_a.png"
    path_b = samples_dir / "pid_drawing_rev_b.png"
    out_dir = tmp_path / "output"
    out_dir.mkdir(parents=True, exist_ok=True)

    boxes = DrawingDiffEngine.compare_revisions(path_a, path_b)
    assert len(boxes) == 3

    # Find boxes by change type
    added_box = next(b for b in boxes if b.change_type == DiffChangeType.ADDED)
    mod_box = next(b for b in boxes if b.change_type == DiffChangeType.MODIFIED)
    rem_box = next(b for b in boxes if b.change_type == DiffChangeType.REMOVED)

    # Compile structured ChangeRegisterItems with engineering classifications
    items = [
        ChangeRegisterItem(
            item_id=added_box.box_id,
            change_type=DiffChangeType.ADDED,
            equipment_tag="RV-402",
            description="Installed new bypass loop around inlet valve V-101 with relief valve RV-402",
            rev_a_state="Not present",
            rev_b_state="Active bypass relief protection",
            safety_risk="HIGH",  # Overpressure protection modification
            box_2d=added_box.box_2d,
        ),
        ChangeRegisterItem(
            item_id=mod_box.box_id,
            change_type=DiffChangeType.MODIFIED,
            equipment_tag="PT-201",
            description="Replaced local analog pressure indicator PI-201 with digital smart transmitter PT-201",
            rev_a_state="Analog indicator PI-201",
            rev_b_state="Digital smart transmitter PT-201 with 4-20mA HART telemetry",
            safety_risk="MEDIUM",  # Instrumentation telemetry upgrade
            box_2d=mod_box.box_2d,
        ),
        ChangeRegisterItem(
            item_id=rem_box.box_id,
            change_type=DiffChangeType.REMOVED,
            equipment_tag="V-102",
            description="Manual drain valve V-102 decommissioned and removed from service",
            rev_a_state="Active drain valve V-102",
            rev_b_state="Decommissioned / flange capped",
            safety_risk="LOW",  # Routine drain removal
            box_2d=rem_box.box_2d,
        ),
        # Adversarial item: tests formula injection sanitization (Threat T-06-04)
        ChangeRegisterItem(
            item_id=99,
            change_type=DiffChangeType.MODIFIED,
            equipment_tag="=cmd|' /C calc'!A0",  # DDE formula injection attempt
            description="@SUM(1+1)*cmd",
            rev_a_state="+100",
            rev_b_state="-200",
            safety_risk="HIGH",
            box_2d=[0.1, 0.1, 0.2, 0.2],
        ),
    ]

    register = ChangeRegister(
        drawing_title="P&ID Unit 04 Hydrogen Header & Vessel V-301",
        rev_a_name=path_a.name,
        rev_b_name=path_b.name,
        items=items,
    )

    # 1. Export Word (.docx) Change Register
    docx_path = out_dir / "pid_revision_change_register.docx"
    exported_docx = ChangeRegisterGenerator.export_word_table(
        register=register,
        docx_path=docx_path,
        allowed_root=tmp_path,
    )

    assert exported_docx.exists()
    assert exported_docx.stat().st_size > 5000

    # Inspect Word deliverable
    doc = docx.Document(str(exported_docx))
    doc_text = " ".join(p.text for p in doc.paragraphs)
    assert "P&ID Revision Verification & Change Register" in doc_text

    assert len(doc.tables) >= 1
    table = doc.tables[0]
    cell_texts = [cell.text for row in table.rows for cell in row.cells]
    assert any("RV-402" in t for t in cell_texts)
    assert any("PT-201" in t for t in cell_texts)
    assert any("V-102" in t for t in cell_texts)
    assert any("HIGH" in t for t in cell_texts)
    assert any("MEDIUM" in t for t in cell_texts)
    assert any("LOW" in t for t in cell_texts)

    # 2. Export CSV Change Register with Formula Injection Sanitization
    csv_path = out_dir / "pid_revision_change_register.csv"
    exported_csv = export_change_register_csv(register=register, csv_path=csv_path)

    assert exported_csv.exists()
    csv_lines = exported_csv.read_text(encoding="utf-8").splitlines()
    assert len(csv_lines) == 1 + len(items)  # Header + items

    # Verify header
    assert csv_lines[0] == "Item #,Change Type,Equipment Tag,Description,Rev A State,Rev B State,Safety Risk"

    # Verify standard items
    csv_full_text = "\n".join(csv_lines)
    assert "RV-402" in csv_full_text
    assert "PT-201" in csv_full_text
    assert "V-102" in csv_full_text

    # Verify Threat T-06-04 Sanitization: No cell may begin with raw '=', '+', '-', or '@'
    with open(exported_csv, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        header = next(reader)
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, cell in enumerate(row):
                if cell:
                    # If the original raw value began with a formula character, it must be escaped with a single quote
                    assert not cell.startswith(("=", "+", "-", "@")), (
                        f"Unsanitized formula character detected in row {row_idx}, col {header[col_idx]}: {cell}"
                    )
