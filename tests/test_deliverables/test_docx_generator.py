"""
tests/test_deliverables/test_docx_generator.py
Unit tests for DocxDeliverableGenerator corporate styling, findings table,
citations appendix, lineage packaging, and path fencing.
"""
import io
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import docx
import pytest

from sovereign.deliverables.docx import DocxDeliverableGenerator
from sovereign.lineage.models import (
    DeliverableInfo,
    LineageMetadata,
    ModelInfo,
    SourceDocumentCitation,
)
from sovereign.rag.citations import Citation
from sovereign.tools.builtin.files import PathSecurityViolation


@pytest.fixture
def sample_findings():
    return [
        {
            "item_num": 1,
            "component": "Distillation Column Overhead Line C-101",
            "observation": "Localized wall thinning detected on outer radius of 90-degree elbow",
            "measured_value": "4.1 mm (min allowable: 5.2 mm)",
            "severity": "Critical",
            "citation": "[Doc: doc-3f8a91b2, p.4#reg-003]",
        },
        {
            "item_num": 2,
            "component": "Reboiler Shell Flange E-204",
            "observation": "Superficial atmospheric oxidation on bolt threads; gasket intact",
            "measured_value": "N/A",
            "severity": "Warning",
            "citation": "[Doc: doc-3f8a91b2, p.7#reg-012]",
        },
        {
            "item_num": 3,
            "component": "Condenser Tubesheet E-102",
            "observation": "Eddy current survey indicates zero tube wall defects exceeding 10%",
            "measured_value": "0.08 mm max pit depth",
            "severity": "Satisfactory",
            "citation": "[Doc: doc-7b12c984, p.2]",
        },
    ]


@pytest.fixture
def sample_citations():
    return [
        Citation(
            document_id="doc-3f8a91b2",
            file_name="UT_Survey_Report_2026_Q3.pdf",
            file_sha256="e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
            page_number=4,
            region_id="reg-003",
            excerpt="Measured wall thickness of elbow E-101 outer radius is 4.1 mm.",
        ),
        Citation(
            document_id="doc-7b12c984",
            file_name="EddyCurrent_Analysis_E102.pdf",
            file_sha256="ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad",
            page_number=2,
            region_id="reg-001",
            excerpt="No pitting observed above 10% wall threshold.",
        ),
    ]


def test_create_approval_note_structure(tmp_path, sample_findings, sample_citations):
    """Test that DocxDeliverableGenerator creates a valid .docx with all expected components."""
    output_path = tmp_path / "approval_note.docx"

    res_path = DocxDeliverableGenerator.create_approval_note(
        output_path=output_path,
        title="Pressure Boundary Integrity Assessment: Column C-101",
        facility="Jurong Island Refinery - Train B",
        equipment_tag="C-101-OH",
        findings=sample_findings,
        citations=sample_citations,
        classification="RESTRICTED",
        author="Senior Reliability Engineer",
    )

    assert res_path == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    # 1. Verify with python-docx library
    doc = docx.Document(output_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any("[ RESTRICTED ]" in p for p in paragraphs)
    assert any("Pressure Boundary Integrity Assessment" in p for p in paragraphs)
    assert any("Jurong Island Refinery" in p for p in paragraphs)
    assert any("C-101-OH" in p for p in paragraphs)

    # Tables verification
    assert len(doc.tables) == 2  # Findings table and References table

    # Table 1: Findings
    findings_tbl = doc.tables[0]
    assert len(findings_tbl.rows) == 1 + len(sample_findings)  # Header + 3 findings
    hdr_cells = [c.text for c in findings_tbl.rows[0].cells]
    assert "Item #" in hdr_cells[0]
    assert "Severity" in hdr_cells[4]

    # Verify critical severity row
    crit_row = findings_tbl.rows[1]
    assert "Critical" in crit_row.cells[4].text
    assert "4.1 mm" in crit_row.cells[3].text

    # Table 2: Source References
    ref_tbl = doc.tables[1]
    assert len(ref_tbl.rows) == 1 + len(sample_citations)
    ref_hdr = [c.text for c in ref_tbl.rows[0].cells]
    assert "Citation Tag" in ref_hdr[0]
    assert "File SHA-256" in ref_hdr[2]


def test_docx_xml_styling_and_shading(tmp_path, sample_findings, sample_citations):
    """Verify raw XML contains corporate colors: navy headers, zebra striping, and severity badges."""
    output_path = tmp_path / "styled_approval.docx"

    DocxDeliverableGenerator.create_approval_note(
        output_path=output_path,
        title="Inspection Report",
        facility="Plant 4",
        equipment_tag="TK-901",
        findings=sample_findings,
        citations=sample_citations,
    )

    with zipfile.ZipFile(output_path, "r") as z:
        assert "word/document.xml" in z.namelist()
        doc_xml = z.read("word/document.xml").decode("utf-8")

        # Check Navy Header (#1F497D)
        assert 'w:fill="1F497D"' in doc_xml

        # Check Zebra striping (#F2F2F2)
        assert 'w:fill="F2F2F2"' in doc_xml

        # Check Severity badges: Critical (#FFD2D2), Warning (#FFF3CD), Satisfactory (#D1E7DD)
        assert 'w:fill="FFD2D2"' in doc_xml
        assert 'w:fill="FFF3CD"' in doc_xml
        assert 'w:fill="D1E7DD"' in doc_xml


def test_docx_lineage_packaging(tmp_path, sample_findings, sample_citations):
    """Verify docProps/custom.xml is injected and sidecar .lineage.json exists with PROV-O data."""
    output_path = tmp_path / "lineage_report.docx"

    custom_metadata = LineageMetadata(
        deliverable=DeliverableInfo(file_name="lineage_report.docx"),
        session_id="session-audit-9999",
        user_id="auditor-01",
        user_role="Lead Auditor",
        model=ModelInfo(model_id="qwen-2.5-coder-32b", weights_sha256="abc123def456"),
        audit_start_seq=10,
        audit_end_seq=15,
        audit_tip_hash="feedface00112233445566778899aabbccddeeff",
    )

    DocxDeliverableGenerator.create_approval_note(
        output_path=output_path,
        title="Lineage Verified Report",
        facility="Plant A",
        equipment_tag="E-100",
        findings=sample_findings,
        citations=sample_citations,
        lineage_metadata=custom_metadata,
    )

    # 1. Sibling sidecar exists and has populated source docs
    sidecar_path = Path(f"{output_path}.lineage.json")
    assert sidecar_path.exists()

    sidecar_data = json.loads(sidecar_path.read_text(encoding="utf-8"))
    assert sidecar_data["session_id"] == "session-audit-9999"
    assert sidecar_data["user_id"] == "auditor-01"
    assert len(sidecar_data["source_documents"]) == len(sample_citations)
    assert sidecar_data["source_documents"][0]["document_id"] == "doc-3f8a91b2"

    # 2. Inspect docProps/custom.xml inside .docx archive
    with zipfile.ZipFile(output_path, "r") as z:
        assert "docProps/custom.xml" in z.namelist()
        custom_xml = z.read("docProps/custom.xml").decode("utf-8")
        assert "sovereign_session_id" in custom_xml
        assert "session-audit-9999" in custom_xml
        assert "sovereign_model_id" in custom_xml
        assert "qwen-2.5-coder-32b" in custom_xml


def test_docx_path_security_fencing(tmp_path, sample_findings):
    """Verify that path traversal outside allowed root or null bytes raise PathSecurityViolation."""
    allowed_root = tmp_path / "workspace"
    allowed_root.mkdir()
    outside_path = tmp_path / "outside" / "evil.docx"

    # Test path escaping root
    with pytest.raises(PathSecurityViolation) as exc_info:
        DocxDeliverableGenerator.create_approval_note(
            output_path=outside_path,
            title="Malicious Report",
            facility="Plant X",
            equipment_tag="X-01",
            findings=sample_findings,
            allowed_root=allowed_root,
        )
    assert "Path traversal detected" in str(exc_info.value)

    # Test null byte injection
    with pytest.raises(PathSecurityViolation) as exc_info_null:
        DocxDeliverableGenerator.create_approval_note(
            output_path=str(allowed_root / "test\x00.docx"),
            title="Null Byte Report",
            facility="Plant X",
            equipment_tag="X-01",
            findings=sample_findings,
            allowed_root=allowed_root,
        )
    assert "Null byte detected" in str(exc_info_null.value)
