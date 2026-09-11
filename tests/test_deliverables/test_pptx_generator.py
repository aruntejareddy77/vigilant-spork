"""
tests/test_deliverables/test_pptx_generator.py
Unit tests for PptxDeliverableGenerator widescreen layouts, native OOXML charts,
tables, bulleted summaries, lineage packaging, and path security fencing.
"""
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.util import Inches
import pytest

from sovereign.deliverables.pptx import (
    PptxDeliverableGenerator,
    SLIDE_HEIGHT_INCHES,
    SLIDE_WIDTH_INCHES,
)
from sovereign.lineage.models import (
    DeliverableInfo,
    LineageMetadata,
    ModelInfo,
)
from sovereign.tools.builtin.files import PathSecurityViolation


@pytest.fixture
def sample_slides_content():
    return [
        {
            "type": "summary",
            "heading": "Executive Summary & Operational Risk",
            "bullet_points": [
                "Overhead line C-101 requires immediate thickness derating or sleeve reinforcement.",
                "Reboiler E-204 and pumps P-101 remain fully compliant with ASME B31.3 allowable stresses.",
                "Next statutory turnaround scheduled for Q1 2027; ultrasonic survey monitoring advised quarterly.",
            ],
        },
        {
            "type": "table",
            "heading": "Critical Component Thickness Summary",
            "headers": ["Tag", "Component", "Nominal (mm)", "Min Req (mm)", "Measured (mm)", "Status"],
            "rows": [
                ["C-101", "Overhead Line", 8.2, 5.2, 4.1, "Critical"],
                ["E-204", "Reboiler Shell", 14.0, 12.0, 12.8, "Satisfactory"],
                ["P-101", "Discharge Elbow", 7.1, 5.0, 6.4, "Satisfactory"],
                ["T-401", "Feed Nozzle", 9.5, 6.8, 7.2, "Satisfactory"],
            ],
        },
    ]


@pytest.fixture
def sample_chart_data():
    return {
        "heading": "Wall Thickness Inspection Comparison",
        "chart_type": "column_clustered",
        "categories": ["C-101", "E-204", "P-101", "T-401"],
        "series": {
            "Measured Actual (mm)": [4.1, 12.8, 6.4, 7.2],
            "Minimum Required (mm)": [5.2, 12.0, 5.0, 6.8],
        },
    }


def test_generate_briefing_deck_structure(tmp_path, sample_slides_content, sample_chart_data):
    """Test generating a 16:9 presentation with title, summary, table, and native chart slides."""
    output_path = tmp_path / "briefing.pptx"

    res_path = PptxDeliverableGenerator.generate_briefing_deck(
        output_path=output_path,
        title="Distillation Column C-101 Integrity Assessment",
        facility="Jurong Island Refinery",
        equipment_tag="C-101",
        slides_content=sample_slides_content,
        chart_data=sample_chart_data,
        classification="RESTRICTED",
        author="Lead Reliability Engineer",
    )

    assert res_path == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    # Inspect with python-pptx
    prs = Presentation(output_path)
    assert len(prs.slides) == 4  # 1 Title + 1 Summary + 1 Table + 1 Chart

    # Verify 16:9 widescreen dimensions
    width_in = prs.slide_width / Inches(1)
    height_in = prs.slide_height / Inches(1)
    assert pytest.approx(width_in, 0.01) == SLIDE_WIDTH_INCHES
    assert pytest.approx(height_in, 0.01) == SLIDE_HEIGHT_INCHES

    # Slide 1: Title slide
    s1 = prs.slides[0]
    s1_texts = [shape.text for shape in s1.shapes if shape.has_text_frame]
    assert any("Distillation Column C-101 Integrity Assessment" in t for t in s1_texts)
    assert any("Jurong Island Refinery" in t for t in s1_texts)

    # Slide 2: Summary slide
    s2 = prs.slides[1]
    s2_texts = [shape.text for shape in s2.shapes if shape.has_text_frame]
    assert any("Executive Summary & Operational Risk" in t for t in s2_texts)
    assert any("Overhead line C-101" in t for t in s2_texts)

    # Slide 3: Table slide
    s3 = prs.slides[2]
    tbl_shapes = [shape for shape in s3.shapes if shape.has_table]
    assert len(tbl_shapes) == 1
    table = tbl_shapes[0].table
    assert len(table.rows) == 5  # 1 header + 4 rows
    assert len(table.columns) == 6

    # Slide 4: Chart slide
    s4 = prs.slides[3]
    chart_shapes = [shape for shape in s4.shapes if shape.has_chart]
    assert len(chart_shapes) == 1
    chart = chart_shapes[0].chart
    assert chart.has_legend is True


def test_pptx_native_ooxml_chart_part_in_archive(tmp_path, sample_slides_content, sample_chart_data):
    """Verify underlying zip contains ppt/charts/chart1.xml with CategoryChartData and values."""
    output_path = tmp_path / "ooxml_chart_check.pptx"

    PptxDeliverableGenerator.generate_briefing_deck(
        output_path=output_path,
        title="Chart Verification Deck",
        facility="Plant 1",
        slides_content=sample_slides_content,
        chart_data=sample_chart_data,
    )

    with zipfile.ZipFile(output_path, "r") as z:
        names = z.namelist()
        # Must have native openxml chart part
        assert "ppt/charts/chart1.xml" in names

        chart_xml = z.read("ppt/charts/chart1.xml").decode("utf-8")
        # Check category names present in OOXML chart
        assert "C-101" in chart_xml
        assert "E-204" in chart_xml
        assert "P-101" in chart_xml
        assert "T-401" in chart_xml

        # Check series data present
        assert "Measured Actual (mm)" in chart_xml
        assert "Minimum Required (mm)" in chart_xml
        assert "4.1" in chart_xml
        assert "12.8" in chart_xml


def test_pptx_lineage_packaging(tmp_path, sample_slides_content, sample_chart_data):
    """Verify docProps/custom.xml and .lineage.json are generated for .pptx deliverables."""
    output_path = tmp_path / "lineage_deck.pptx"

    meta = LineageMetadata(
        deliverable=DeliverableInfo(file_name="lineage_deck.pptx"),
        session_id="session-deck-777",
        user_id="chief_engineer",
        user_role="Chief Mechanical Engineer",
        model=ModelInfo(model_id="qwen-2.5-coder-32b", weights_sha256="deadbeef9988"),
        audit_start_seq=2,
        audit_end_seq=6,
        audit_tip_hash="abcdef00112233445566778899aabbccddeeff",
    )

    PptxDeliverableGenerator.generate_briefing_deck(
        output_path=output_path,
        title="Lineage Deck",
        facility="Offshore Platform Alpha",
        slides_content=sample_slides_content,
        chart_data=sample_chart_data,
        lineage_metadata=meta,
    )

    # Check sidecar manifest
    sidecar = Path(f"{output_path}.lineage.json")
    assert sidecar.exists()
    sidecar_data = json.loads(sidecar.read_text(encoding="utf-8"))
    assert sidecar_data["session_id"] == "session-deck-777"
    assert sidecar_data["user_id"] == "chief_engineer"

    # Check docProps/custom.xml in pptx zip
    with zipfile.ZipFile(output_path, "r") as z:
        assert "docProps/custom.xml" in z.namelist()
        custom_xml = z.read("docProps/custom.xml").decode("utf-8")
        assert "session-deck-777" in custom_xml
        assert "qwen-2.5-coder-32b" in custom_xml


def test_pptx_path_security_fencing(tmp_path, sample_slides_content):
    """Verify path traversal outside allowed root and null bytes raise PathSecurityViolation."""
    allowed_root = tmp_path / "safe_dir"
    allowed_root.mkdir()
    forbidden_path = tmp_path / "forbidden.pptx"

    with pytest.raises(PathSecurityViolation):
        PptxDeliverableGenerator.generate_briefing_deck(
            output_path=forbidden_path,
            title="Breach",
            facility="Refinery",
            slides_content=sample_slides_content,
            allowed_root=allowed_root,
        )

    with pytest.raises(PathSecurityViolation):
        PptxDeliverableGenerator.generate_briefing_deck(
            output_path=str(allowed_root / "deck\x00.pptx"),
            title="NullByte",
            facility="Refinery",
            slides_content=sample_slides_content,
            allowed_root=allowed_root,
        )
