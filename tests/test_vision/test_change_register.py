"""
tests/test_vision/test_change_register.py
Unit and Integration Tests for ChangeRegisterGenerator and diff_drawings Tool.
Validates Layer 2 semantic VLM reasoning, ChangeRegister data structures,
Word table export with lineage sidecars, and ToolRegistry RBAC dispatching.
Complies with MMOD-03, DELV-01, and ASVS V5.1/V10.3.
"""
from pathlib import Path
import pytest
from PIL import Image, ImageDraw
import docx

from sovereign.models.adapters.mock import MockModelAdapter
from sovereign.models.manifest import load_manifest
from sovereign.tools.builtin.files import PathSecurityViolation
from sovereign.tools.registry import ToolExecutionError, ToolRegistry
from sovereign.tools.schemas import UserRole
from sovereign.vision.diff import DiffBoundingBox, DiffChangeType, DrawingDiffEngine
from sovereign.vision.register import (
    ChangeRegister,
    ChangeRegisterGenerator,
    ChangeRegisterItem,
)
from sovereign.vision.tools import (
    DIFF_DRAWINGS_TOOL_DEF,
    handle_diff_drawings,
    register_vision_tools,
)

FIXTURES_DIR = Path(__file__).parent.parent / "fixtures" / "models"


@pytest.fixture
def mock_adapter() -> MockModelAdapter:
    """Provides a deterministic MockModelAdapter for vision reasoning tests."""
    manifest = load_manifest(FIXTURES_DIR / "mock_default.yaml")
    return MockModelAdapter(manifest)


@pytest.fixture
def sample_drawings(tmp_path: Path):
    """Generates synthetic P&ID drawing pair for change register testing."""
    width, height = 800, 600

    # Rev A
    im_a = Image.new("RGB", (width, height), (255, 255, 255))
    draw_a = ImageDraw.Draw(im_a)
    draw_a.line([(100, 300), (700, 300)], fill=(0, 0, 0), width=6)
    draw_a.rectangle([230, 280, 270, 320], fill=(0, 0, 0))  # Valve V-101 (unchanged)
    draw_a.rectangle([580, 280, 620, 320], fill=(0, 0, 0))  # Valve V-102 (removed)
    draw_a.rectangle([380, 90, 420, 130], fill=(0, 0, 0))   # Tag (modified)

    path_a = tmp_path / "pid_rev_a.png"
    im_a.save(path_a)

    # Rev B
    im_b = Image.new("RGB", (width, height), (255, 255, 255))
    draw_b = ImageDraw.Draw(im_b)
    draw_b.line([(100, 300), (700, 300)], fill=(0, 0, 0), width=6)
    draw_b.rectangle([230, 280, 270, 320], fill=(0, 0, 0))  # Valve V-101
    draw_b.rectangle([130, 90, 170, 130], fill=(0, 0, 0))   # Pump P-101 (added)
    draw_b.ellipse([360, 70, 440, 150], fill=(120, 120, 120))  # Tag (modified)

    path_b = tmp_path / "pid_rev_b.png"
    im_b.save(path_b)

    return path_a, path_b


class TestChangeRegisterModels:
    """Unit tests for ChangeRegister and ChangeRegisterItem data models."""

    def test_item_and_register_serialization(self):
        """Validates JSON round-trip serialization and field types."""
        item = ChangeRegisterItem(
            item_id=1,
            change_type=DiffChangeType.ADDED,
            equipment_tag="CV-103",
            description="Installed high-pressure bypass control valve",
            rev_a_state="Line absent",
            rev_b_state="Installed and commissioned",
            safety_risk="HIGH",
            box_2d=[0.1, 0.2, 0.3, 0.4],
        )
        data = item.model_dump()
        assert data["equipment_tag"] == "CV-103"
        assert data["safety_risk"] == "HIGH"

        reg = ChangeRegister(
            drawing_title="P&ID Unit 100",
            rev_a_name="Rev_A.png",
            rev_b_name="Rev_B.png",
            items=[item],
        )
        reg_json = reg.model_dump_json()
        assert "CV-103" in reg_json

        rehydrated = ChangeRegister.model_validate_json(reg_json)
        assert len(rehydrated.items) == 1
        assert rehydrated.items[0].equipment_tag == "CV-103"


class TestChangeRegisterGenerator:
    """Unit tests for ChangeRegisterGenerator logic and multi-format exports."""

    def test_generate_register_with_vlm_adapter(self, sample_drawings, mock_adapter):
        """Validates VLM reasoning integration extracting semantic tags and descriptions."""
        path_a, path_b = sample_drawings
        diff_boxes = DrawingDiffEngine.compare_revisions(path_a, path_b)
        assert len(diff_boxes) >= 2

        register = ChangeRegisterGenerator.generate_register(
            diff_boxes=diff_boxes,
            model_adapter=mock_adapter,
            img_a_path=path_a,
            img_b_path=path_b,
            drawing_title="P&ID-SYS-101 Hydrocarbon Feed",
        )

        assert register.drawing_title == "P&ID-SYS-101 Hydrocarbon Feed"
        assert register.rev_a_name == "pid_rev_a.png"
        assert register.rev_b_name == "pid_rev_b.png"
        assert len(register.items) == len(diff_boxes)

        for item in register.items:
            assert item.item_id > 0
            assert item.equipment_tag != ""
            assert item.description != ""
            assert item.safety_risk in ["LOW", "MEDIUM", "HIGH"]
            assert len(item.box_2d) == 4

    def test_generate_register_fallback_without_vlm(self, sample_drawings):
        """Verifies deterministic fallback descriptions when model adapter is None."""
        path_a, path_b = sample_drawings
        diff_boxes = DrawingDiffEngine.compare_revisions(path_a, path_b)

        register = ChangeRegisterGenerator.generate_register(
            diff_boxes=diff_boxes,
            model_adapter=None,
            img_a_path=path_a,
            img_b_path=path_b,
        )

        assert len(register.items) == len(diff_boxes)
        assert "TAG-01" in register.items[0].equipment_tag

    def test_export_markdown(self, sample_drawings):
        """Verifies markdown table formatting."""
        item = ChangeRegisterItem(
            item_id=1,
            change_type=DiffChangeType.ADDED,
            equipment_tag="P-101",
            description="Installed centrifugal feed pump",
            rev_a_state="Open header",
            rev_b_state="Active pump station",
            safety_risk="MEDIUM",
            box_2d=[0.1, 0.1, 0.2, 0.2],
        )
        reg = ChangeRegister(
            drawing_title="Cooling Water Schematic",
            rev_a_name="rev_01.png",
            rev_b_name="rev_02.png",
            items=[item],
        )

        md = ChangeRegisterGenerator.export_markdown(reg)
        assert "# Engineering Change Register: Cooling Water Schematic" in md
        assert "| Item # | Change Type | Equipment Tag | Description |" in md
        assert "| 1 | ADDED | `P-101` | Installed centrifugal feed pump |" in md

    def test_export_word_table(self, sample_drawings, tmp_path: Path):
        """Verifies native Word table deliverable generation with lineage metadata sidecar."""
        item1 = ChangeRegisterItem(
            item_id=1,
            change_type=DiffChangeType.ADDED,
            equipment_tag="P-101",
            description="Added feed pump P-101",
            rev_a_state="Not present",
            rev_b_state="Installed",
            safety_risk="HIGH",
            box_2d=[0.1, 0.1, 0.2, 0.2],
        )
        item2 = ChangeRegisterItem(
            item_id=2,
            change_type=DiffChangeType.REMOVED,
            equipment_tag="V-102",
            description="Decommissioned gate valve V-102",
            rev_a_state="In line",
            rev_b_state="Removed",
            safety_risk="LOW",
            box_2d=[0.5, 0.5, 0.6, 0.6],
        )
        reg = ChangeRegister(
            drawing_title="Main Process P&ID",
            rev_a_name="rev_a.png",
            rev_b_name="rev_b.png",
            items=[item1, item2],
        )

        docx_path = tmp_path / "deliverables" / "Change_Register.docx"
        result_path = ChangeRegisterGenerator.export_word_table(reg, docx_path)

        assert result_path.exists()
        assert result_path == docx_path.resolve()

        # Check document structure
        doc = docx.Document(result_path)
        assert len(doc.tables) >= 1
        table = doc.tables[0]
        # 1 header row + 2 data rows = 3 rows
        assert len(table.rows) == 3

        # Sibling lineage manifest
        sidecar = Path(f"{result_path}.lineage.json")
        assert sidecar.exists()
        assert "sovereign-vision-register-v1" in sidecar.read_text(encoding="utf-8")


class TestVisionTools:
    """Integration tests for the diff_drawings tool registered in ToolRegistry."""

    def test_tool_definition_and_clearance(self):
        """Validates tool schema and clearance requirements."""
        assert DIFF_DRAWINGS_TOOL_DEF.name == "diff_drawings"
        assert DIFF_DRAWINGS_TOOL_DEF.required_role == UserRole.ENGINEER
        assert "rev_a_path" in DIFF_DRAWINGS_TOOL_DEF.parameters_schema["required"]
        assert "rev_b_path" in DIFF_DRAWINGS_TOOL_DEF.parameters_schema["required"]
        assert "output_overlay_path" in DIFF_DRAWINGS_TOOL_DEF.parameters_schema["required"]

    def test_diff_drawings_tool_execution(self, sample_drawings, mock_adapter, tmp_path: Path):
        """Executes diff_drawings tool via ToolRegistry and verifies deliverables."""
        path_a, path_b = sample_drawings
        registry = ToolRegistry()
        register_vision_tools(registry, model_adapter=mock_adapter, workspace_root=tmp_path)

        out_overlay = tmp_path / "overlay.png"
        out_report = tmp_path / "change_report.docx"

        result = registry.execute_tool(
            name="diff_drawings",
            arguments={
                "rev_a_path": str(path_a),
                "rev_b_path": str(path_b),
                "output_overlay_path": str(out_overlay),
                "output_report_path": str(out_report),
                "drawing_title": "Primary Crude Distillation P&ID",
            },
            user_role=UserRole.ENGINEER,
        )

        assert result["status"] == "success"
        assert result["total_changes"] >= 2
        assert out_overlay.exists()
        assert out_report.exists()
        assert len(result["items"]) >= 2
        assert "Primary Crude Distillation P&ID" in result["markdown"]

    def test_diff_drawings_rbac_enforcement(self, sample_drawings, tmp_path: Path):
        """Users without ENGINEER clearance cannot execute diff_drawings."""
        path_a, path_b = sample_drawings
        registry = ToolRegistry()
        register_vision_tools(registry, workspace_root=tmp_path)

        with pytest.raises(PermissionError, match="does not have sufficient clearance"):
            registry.execute_tool(
                name="diff_drawings",
                arguments={
                    "rev_a_path": str(path_a),
                    "rev_b_path": str(path_b),
                    "output_overlay_path": str(tmp_path / "out.png"),
                },
                user_role=UserRole.AUDITOR,
            )

    def test_diff_drawings_path_traversal_fencing(self, sample_drawings, tmp_path: Path):
        """Path traversal outside workspace root raises PathSecurityViolation."""
        path_a, path_b = sample_drawings
        workspace = tmp_path / "safe_workspace"
        workspace.mkdir()

        registry = ToolRegistry()
        register_vision_tools(registry, workspace_root=workspace)

        with pytest.raises((PathSecurityViolation, ToolExecutionError), match="Path traversal detected"):
            registry.execute_tool(
                name="diff_drawings",
                arguments={
                    "rev_a_path": str(path_a),
                    "rev_b_path": str(path_b),
                    "output_overlay_path": str(tmp_path / "escaped_overlay.png"),
                },
                user_role=UserRole.ENGINEER,
            )

        with pytest.raises(PathSecurityViolation, match="Path traversal detected"):
            handle_diff_drawings(
                rev_a_path=str(path_a),
                rev_b_path=str(path_b),
                output_overlay_path=str(tmp_path / "escaped_overlay.png"),
                workspace_root=workspace,
            )
