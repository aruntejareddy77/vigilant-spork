"""
tests/test_deliverables/test_xlsx_generator.py
Unit tests for XlsxDeliverableGenerator live formulas, 3-tier visual hierarchy,
formula injection defense, lineage packaging, and path fencing.
"""
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest

from sovereign.deliverables.xlsx import XlsxDeliverableGenerator, sanitize_cell_text
from sovereign.lineage.models import (
    DeliverableInfo,
    LineageMetadata,
    ModelInfo,
)
from sovereign.tools.builtin.files import PathSecurityViolation


@pytest.fixture
def sample_asme_inputs():
    return [
        {"name": "Design Internal Pressure", "symbol": "P", "value": 2.5, "unit": "MPa", "notes": "Design Spec DS-101"},
        {"name": "Pipe Outside Diameter", "symbol": "D", "value": 168.3, "unit": "mm", "notes": "DN150 / 6-inch pipe"},
        {"name": "Basic Allowable Stress", "symbol": "S", "value": 138.0, "unit": "MPa", "notes": "ASTM A106 Gr. B @ 150C"},
        {"name": "Joint Quality Factor", "symbol": "E", "value": 1.0, "unit": "-", "notes": "Seamless pipe Table A-1B"},
        {"name": "Weld Joint Strength Factor", "symbol": "W", "value": 1.0, "unit": "-", "notes": "Non-welded / seamless"},
        {"name": "Wall Thickness Coefficient", "symbol": "Y", "value": 0.4, "unit": "-", "notes": "Ferritic steel < 482C"},
        {"name": "Corrosion Allowance", "symbol": "c", "value": 3.0, "unit": "mm", "notes": "Sour service margin"},
        {"name": "Mill Undervalue Tolerance", "symbol": "MT", "value": 0.125, "unit": "-", "notes": "12.5% ASTM tolerance"},
        {"name": "Measured Actual Thickness", "symbol": "t_act", "value": 5.4, "unit": "mm", "notes": "UT survey 2026-Q3"},
    ]


@pytest.fixture
def sample_asme_calcs():
    # Note: Row 7: P, Row 8: D, Row 9: S, Row 10: E, Row 11: W, Row 12: Y, Row 13: c, Row 14: MT, Row 15: t_act
    # Calcs start at row 18:
    # Row 19: Pressure design thickness t = (P*D)/(2*(S*E*W + P*Y)) -> (C7*C8)/(2*(C9*C10*C11 + C7*C12))
    # Row 20: Min required thickness t_m = t + c -> C19 + C13
    # Row 21: Min nominal thickness t_nom_req = t_m / (1 - MT) -> C20 / (1 - C14)
    return [
        {
            "name": "Pressure Design Thickness",
            "symbol": "t",
            "formula_excel": "=C7*C8/(2*(C9*C10*C11+C7*C12))",
            "unit": "mm",
            "clause": "ASME B31.3 Eq. 3a",
        },
        {
            "name": "Minimum Required Thickness",
            "symbol": "t_m",
            "formula_excel": "=C19+C13",
            "unit": "mm",
            "clause": "ASME B31.3 Clause 304.1.1",
        },
        {
            "name": "Minimum Nominal Thickness (incl. Mill Tol)",
            "symbol": "t_nom_req",
            "formula_excel": "=C20/(1-C14)",
            "unit": "mm",
            "clause": "ASME B31.3 Clause 304.1.2",
        },
    ]


@pytest.fixture
def sample_status_formula():
    # Compare measured thickness (C15) against minimum required (C20)
    return '=IF(C15>=C20, "PASS - SAFE", "FAIL - UNSAFE")'


def test_generate_calculation_workbook_structure(tmp_path, sample_asme_inputs, sample_asme_calcs, sample_status_formula):
    """Test generating a workbook and verifying structure, formulas, and cells with openpyxl."""
    output_path = tmp_path / "asme_calc.xlsx"

    res_path = XlsxDeliverableGenerator.generate_calculation_workbook(
        output_path=output_path,
        title="Piping Pressure Design Thickness Verification",
        standard_ref="ASME B31.3-2022 Process Piping Clause 304.1.2",
        inputs=sample_asme_inputs,
        calculations=sample_asme_calcs,
        conclusion_formula=sample_status_formula,
        facility="Jurong Island Train B",
        equipment_tag="Line-01-P-101",
    )

    assert res_path == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    # Load with openpyxl
    wb = openpyxl.load_workbook(output_path, data_only=False)
    assert "Engineering Calculation" in wb.sheetnames
    ws = wb["Engineering Calculation"]

    # Title & Subtitle checks
    assert "Piping Pressure Design Thickness Verification" in ws["A1"].value
    assert "ASME B31.3-2022" in ws["A2"].value

    # Check input values
    assert ws["C7"].value == 2.5
    assert ws["C8"].value == 168.3

    # Check live calculation formula strings
    assert ws["C19"].value == "=C7*C8/(2*(C9*C10*C11+C7*C12))"
    assert ws["C20"].value == "=C19+C13"
    assert ws["C21"].value == "=C20/(1-C14)"

    # Check status formula string
    assert ws["C25"].value == '=IF(C15>=C20, "PASS - SAFE", "FAIL - UNSAFE")'

    # Check 3-tier colors
    # Tier 1 Input: D9E1F2
    assert "D9E1F2" in str(ws["C7"].fill.start_color.rgb).upper()

    # Tier 2 Calc: FFF2CC
    assert "FFF2CC" in str(ws["C19"].fill.start_color.rgb).upper()

    # Tier 3 Status: E2EFDA
    assert "E2EFDA" in str(ws["C25"].fill.start_color.rgb).upper()


def test_xlsx_raw_xml_contains_formula_tags(tmp_path, sample_asme_inputs, sample_asme_calcs, sample_status_formula):
    """Verify underlying worksheet XML contains native <f> formula elements."""
    output_path = tmp_path / "raw_xml_check.xlsx"

    XlsxDeliverableGenerator.generate_calculation_workbook(
        output_path=output_path,
        title="XML Formula Tag Verification",
        standard_ref="ASME B31.3",
        inputs=sample_asme_inputs,
        calculations=sample_asme_calcs,
        conclusion_formula=sample_status_formula,
    )

    with zipfile.ZipFile(output_path, "r") as z:
        assert "xl/worksheets/sheet1.xml" in z.namelist()
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")

        # Must have native openxml <f> formula tags
        assert "<f>" in sheet_xml
        assert "</f>" in sheet_xml
        # Verify formula content is inside the XML
        assert "C7*C8/(2*(C9*C10*C11+C7*C12))" in sheet_xml
        assert "C19+C13" in sheet_xml
        assert 'IF(C15&gt;=C20' in sheet_xml or 'IF(C15>=C20' in sheet_xml


def test_formula_injection_sanitizer():
    """Verify that sanitize_cell_text escapes DDE formula trigger chars with a leading apostrophe."""
    # Trigger chars: =, +, -, @
    assert sanitize_cell_text("=cmd|' /C calc'!A0") == "'=cmd|' /C calc'!A0"
    assert sanitize_cell_text("+123456") == "'+123456"
    assert sanitize_cell_text("-10.5mm") == "'-10.5mm"
    assert sanitize_cell_text("@SUM(A1:A10)") == "'@SUM(A1:A10)"

    # Safe text and numbers must not be modified
    assert sanitize_cell_text("Normal text") == "Normal text"
    assert sanitize_cell_text("DN150 Pipe") == "DN150 Pipe"
    assert sanitize_cell_text(123.45) == 123.45
    assert sanitize_cell_text(42) == 42
    assert sanitize_cell_text("") == ""
    assert sanitize_cell_text(None) is None


def test_formula_injection_in_workbook_cells(tmp_path):
    """Verify that malicious inputs starting with formula triggers are escaped in the workbook."""
    output_path = tmp_path / "sanitized_workbook.xlsx"

    malicious_inputs = [
        {"name": "=1+1", "symbol": "+EVIL", "value": "@SUM(1,2)", "unit": "-calc", "notes": "=cmd|' /C calc'"},
    ]
    calcs = [
        {"name": "Legit Step", "symbol": "R", "formula_excel": "=10*2", "unit": "mm", "clause": "Standard"},
    ]

    XlsxDeliverableGenerator.generate_calculation_workbook(
        output_path=output_path,
        title="=InjectionTitle",
        standard_ref="+InjectionStandard",
        inputs=malicious_inputs,
        calculations=calcs,
        conclusion_formula='=IF(C7>0, "PASS", "FAIL")',
    )

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Engineering Calculation"]

    # Title & standard escaped
    assert ws["A1"].value.startswith("'=")
    assert ws["A2"].value.startswith("'+") or "'+InjectionStandard" in ws["A2"].value

    # Malicious inputs escaped
    assert ws["A7"].value.startswith("'=")
    assert ws["B7"].value.startswith("'+")
    assert ws["C7"].value.startswith("'@")
    assert ws["D7"].value.startswith("'-")
    assert ws["E7"].value.startswith("'=")

    # Authorized calculation formula remains active formula
    assert ws["C11"].value == "=10*2"


def test_xlsx_lineage_packaging(tmp_path, sample_asme_inputs, sample_asme_calcs, sample_status_formula):
    """Verify docProps/custom.xml and .lineage.json are generated for .xlsx deliverables."""
    output_path = tmp_path / "lineage_calc.xlsx"

    meta = LineageMetadata(
        deliverable=DeliverableInfo(file_name="lineage_calc.xlsx"),
        session_id="session-calc-001",
        user_id="lead_eng",
        user_role="Senior Process Engineer",
        model=ModelInfo(model_id="deepseek-r1-distill", weights_sha256="1234567890abcdef"),
        audit_start_seq=4,
        audit_end_seq=8,
        audit_tip_hash="abcdef0123456789abcdef0123456789",
    )

    XlsxDeliverableGenerator.generate_calculation_workbook(
        output_path=output_path,
        title="Lineage ASME Sheet",
        standard_ref="ASME B31.3",
        inputs=sample_asme_inputs,
        calculations=sample_asme_calcs,
        conclusion_formula=sample_status_formula,
        lineage_metadata=meta,
    )

    sidecar = Path(f"{output_path}.lineage.json")
    assert sidecar.exists()
    sidecar_data = json.loads(sidecar.read_text(encoding="utf-8"))
    assert sidecar_data["session_id"] == "session-calc-001"
    assert sidecar_data["user_id"] == "lead_eng"

    with zipfile.ZipFile(output_path, "r") as z:
        assert "docProps/custom.xml" in z.namelist()
        custom_xml = z.read("docProps/custom.xml").decode("utf-8")
        assert "session-calc-001" in custom_xml
        assert "deepseek-r1-distill" in custom_xml


def test_xlsx_path_security_fencing(tmp_path, sample_asme_inputs, sample_asme_calcs, sample_status_formula):
    """Verify path traversal outside allowed root or null bytes raise PathSecurityViolation."""
    allowed_root = tmp_path / "sandbox"
    allowed_root.mkdir()
    outside_path = tmp_path / "forbidden.xlsx"

    with pytest.raises(PathSecurityViolation):
        XlsxDeliverableGenerator.generate_calculation_workbook(
            output_path=outside_path,
            title="Breach",
            standard_ref="ASME",
            inputs=sample_asme_inputs,
            calculations=sample_asme_calcs,
            conclusion_formula=sample_status_formula,
            allowed_root=allowed_root,
        )

    with pytest.raises(PathSecurityViolation):
        XlsxDeliverableGenerator.generate_calculation_workbook(
            output_path=str(allowed_root / "calc\x00.xlsx"),
            title="NullByte",
            standard_ref="ASME",
            inputs=sample_asme_inputs,
            calculations=sample_asme_calcs,
            conclusion_formula=sample_status_formula,
            allowed_root=allowed_root,
        )
