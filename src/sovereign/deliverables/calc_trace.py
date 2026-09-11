"""
src/sovereign/deliverables/calc_trace.py
Engineering calculation trace exporter formatting step-by-step mathematical derivations
(ASME B31.3, API 570) to Markdown, canonical JSON, Word sections, and live Excel workbooks.
Complies with DELV-05.
"""
from datetime import datetime, timezone
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import uuid

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field

from sovereign.deliverables.docx import (
    COLOR_BODY_CHARCOAL,
    COLOR_PRIMARY_NAVY,
    COLOR_SECONDARY_SLATE,
    COLOR_WHITE,
    HEX_BADGE_CRITICAL_BG,
    HEX_BADGE_SATISFACTORY_BG,
    HEX_PRIMARY_NAVY,
    HEX_ZEBRA_FILL,
    set_cell_margins,
    set_cell_shading,
    validate_output_path,
)
from sovereign.deliverables.xlsx import XlsxDeliverableGenerator
from sovereign.lineage.models import LineageMetadata


class CalculationParameter(BaseModel):
    """Specification of an engineering input parameter or design constant."""
    model_config = ConfigDict(extra="allow")

    name: str
    symbol: str
    value: Union[float, int, str]
    unit: str
    description: Optional[str] = None
    clause: Optional[str] = None
    citation: Optional[str] = None


class CalculationStep(BaseModel):
    """Specification of an intermediate mathematical derivation step."""
    model_config = ConfigDict(extra="allow")

    name: str
    symbol: str
    formula_display: str
    formula_latex: Optional[str] = None
    formula_excel: str
    intermediate_substitution: str
    result_value: float
    unit: str
    clause: Optional[str] = None


class CalculationTrace(BaseModel):
    """Complete traceable engineering derivation package."""
    model_config = ConfigDict(extra="allow")

    title: str
    governing_standard: str
    facility: Optional[str] = None
    equipment_tag: Optional[str] = None
    parameters: List[CalculationParameter] = Field(default_factory=list)
    steps: List[CalculationStep] = Field(default_factory=list)
    conclusion_formula: str
    status_condition: str
    is_compliant: bool = True
    safety_margin: float = 0.0
    summary: Optional[str] = None


class CalculationTraceExporter:
    """
    Exports structured engineering calculation traces across multiple native formats:
    Markdown, canonical JSON, Word appendix, and live Excel workbooks.
    """

    def __init__(self, trace: CalculationTrace):
        self.trace = trace

    def to_markdown(self) -> str:
        """Renders the calculation trace to clean GitHub-flavored Markdown."""
        lines = [
            f"# {self.trace.title}",
            f"**Governing Standard:** {self.trace.governing_standard}  ",
        ]
        if self.trace.facility or self.trace.equipment_tag:
            lines.append(f"**Facility:** {self.trace.facility or 'N/A'} | **Equipment Tag:** {self.trace.equipment_tag or 'N/A'}  ")
        lines.append("")

        if self.trace.summary:
            lines.append(f"*{self.trace.summary}*")
            lines.append("")

        # Section 1: Parameters Table
        lines.append("## 1. Input Parameters & Design Conditions")
        lines.append("| Parameter | Symbol | Value | Unit | Governing Clause / Reference |")
        lines.append("|:---|:---:|:---:|:---:|:---|")
        for p in self.trace.parameters:
            ref = p.clause or p.citation or "-"
            lines.append(f"| {p.name} | **{p.symbol}** | `{p.value}` | {p.unit} | {ref} |")
        lines.append("")

        # Section 2: Mathematical Derivations
        lines.append("## 2. Step-by-Step Mathematical Derivations")
        for idx, s in enumerate(self.trace.steps, start=1):
            lines.append(f"### Step 2.{idx}: {s.name} ({s.symbol})")
            if s.clause:
                lines.append(f"**Clause Reference:** {s.clause}")
            lines.append(f"- **Governing Formula:** `{s.formula_display}`")
            if s.formula_latex:
                lines.append(f"  $$\n  {s.formula_latex}\n  $$")
            lines.append(f"- **Numerical Substitution:** `{s.intermediate_substitution}`")
            lines.append(f"- **Calculated Result:** **`{s.result_value:.4f} {s.unit}`**")
            lines.append(f"- **Active Spreadsheet Formula:** `{s.formula_excel}`")
            lines.append("")

        # Section 3: Verification & Compliance
        status_badge = "**PASS - SAFE**" if self.trace.is_compliant else "**FAIL - UNSAFE**"
        lines.append("## 3. Compliance & Verification Status")
        lines.append(f"- **Status Evaluation:** {status_badge}")
        lines.append(f"- **Condition Verified:** `{self.trace.status_condition}`")
        lines.append(f"- **Safety Margin:** `{self.trace.safety_margin:+.3f} mm`")
        lines.append(f"- **Live Logic:** `{self.trace.conclusion_formula}`")

        return "\n".join(lines)

    def to_json(self, indent: int = 2) -> str:
        """Exports canonical JSON representation of the calculation trace."""
        return self.trace.model_dump_json(indent=indent)

    def to_excel_workbook(
        self,
        output_path: Union[str, Path],
        lineage_metadata: Optional[LineageMetadata] = None,
        allowed_root: Optional[Union[str, Path]] = None,
    ) -> Path:
        """
        Exports the calculation trace as a native live Excel workbook via XlsxDeliverableGenerator.
        """
        # Convert parameters to inputs format
        excel_inputs = [
            {
                "name": p.name,
                "symbol": p.symbol,
                "value": p.value,
                "unit": p.unit,
                "notes": p.clause or p.citation or "",
            }
            for p in self.trace.parameters
        ]

        # Convert steps to calcs format
        excel_calcs = [
            {
                "name": s.name,
                "symbol": s.symbol,
                "formula_excel": s.formula_excel,
                "unit": s.unit,
                "clause": s.clause or "",
            }
            for s in self.trace.steps
        ]

        return XlsxDeliverableGenerator.generate_calculation_workbook(
            output_path=output_path,
            title=self.trace.title,
            standard_ref=self.trace.governing_standard,
            inputs=excel_inputs,
            calculations=excel_calcs,
            conclusion_formula=self.trace.conclusion_formula,
            lineage_metadata=lineage_metadata,
            facility=self.trace.facility,
            equipment_tag=self.trace.equipment_tag,
            allowed_root=allowed_root,
        )

    def append_to_docx(self, doc: docx.Document) -> None:
        """Appends a complete engineering derivation section to a python-docx Document."""
        # Heading 1
        h1 = doc.add_paragraph()
        r1 = h1.add_run(f"Engineering Calculation Trace: {self.trace.title}")
        r1.font.name = "Calibri"
        r1.font.size = Pt(14)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_PRIMARY_NAVY
        h1.paragraph_format.space_after = Pt(4)

        # Subheading
        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run(f"Governing Standard: {self.trace.governing_standard}")
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(10)
        r_sub.font.italic = True
        r_sub.font.color.rgb = COLOR_SECONDARY_SLATE

        # 1. Parameters Table
        p_param = doc.add_paragraph("Input Parameters & Operating Variables")
        p_param.style.font.bold = True
        p_param.paragraph_format.space_before = Pt(8)
        p_param.paragraph_format.space_after = Pt(4)

        headers = ["Parameter", "Symbol", "Value", "Unit", "Clause / Reference"]
        tbl = doc.add_table(rows=1 + len(self.trace.parameters), cols=5)
        tbl.style = "Table Grid"

        # Table Header
        hdr_row = tbl.rows[0]
        for idx, h_text in enumerate(headers):
            cell = hdr_row.cells[idx]
            set_cell_shading(cell, HEX_PRIMARY_NAVY)
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_WHITE

        # Table Rows
        for r_idx, p in enumerate(self.trace.parameters, start=1):
            row = tbl.rows[r_idx]
            is_odd = (r_idx % 2 == 1)
            row_fill = HEX_ZEBRA_FILL if is_odd else "FFFFFF"

            vals = [p.name, p.symbol, str(p.value), p.unit, p.clause or p.citation or "-"]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                set_cell_shading(cell, row_fill)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                cp = cell.paragraphs[0]
                run = cp.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_BODY_CHARCOAL
                if c_idx == 1:
                    run.font.bold = True

        # 2. Steps
        p_steps = doc.add_paragraph("Step-by-Step Mathematical Derivations")
        p_steps.style.font.bold = True
        p_steps.paragraph_format.space_before = Pt(12)
        p_steps.paragraph_format.space_after = Pt(4)

        for s_idx, s in enumerate(self.trace.steps, start=1):
            p_step = doc.add_paragraph()
            r_st = p_step.add_run(f"Step {s_idx}: {s.name} ({s.symbol})\n")
            r_st.font.bold = True
            r_st.font.size = Pt(10)
            r_st.font.color.rgb = COLOR_PRIMARY_NAVY

            p_step.add_run(f"  • Formula: {s.formula_display}\n")
            p_step.add_run(f"  • Substitution: {s.intermediate_substitution}\n")
            r_res = p_step.add_run(f"  • Result: {s.result_value:.4f} {s.unit}")
            r_res.font.bold = True
            if s.clause:
                p_step.add_run(f"  (Ref: {s.clause})")
            p_step.paragraph_format.space_after = Pt(6)

        # 3. Compliance Callout
        p_comp = doc.add_paragraph()
        status_str = "PASS - COMPLIANT" if self.trace.is_compliant else "FAIL - NON-COMPLIANT"
        r_c = p_comp.add_run(f"Verification Assessment: {status_str}\n")
        r_c.font.bold = True
        r_c.font.size = Pt(11)
        r_c.font.color.rgb = RGBColor(0x0F, 0x51, 0x32) if self.trace.is_compliant else RGBColor(0x90, 0x00, 0x00)

        p_comp.add_run(f"Condition: {self.trace.status_condition} | Margin: {self.trace.safety_margin:+.3f} mm")
        p_comp.paragraph_format.space_after = Pt(12)

    @classmethod
    def create_asme_b313_trace(
        cls,
        pressure_mpa: float = 2.5,
        diameter_mm: float = 168.3,
        stress_mpa: float = 138.0,
        quality_factor: float = 1.0,
        weld_factor: float = 1.0,
        y_coeff: float = 0.4,
        corrosion_allowance_mm: float = 3.0,
        mill_tolerance: float = 0.125,
        actual_thickness_mm: float = 5.4,
        facility: str = "Jurong Island Refinery",
        equipment_tag: str = "Line-01-P-101",
    ) -> CalculationTrace:
        """
        Factory creating an ASME B31.3 Clause 304.1.2 pressure design calculation trace.
        Computes:
        t = (P * D) / (2 * (S * E * W + P * Y))
        t_m = t + c
        t_nom_req = t_m / (1 - MT)
        margin = actual_thickness - t_m
        """
        P = pressure_mpa
        D = diameter_mm
        S = stress_mpa
        E = quality_factor
        W = weld_factor
        Y = y_coeff
        c = corrosion_allowance_mm
        MT = mill_tolerance
        t_act = actual_thickness_mm

        # Math derivations
        denom = 2.0 * (S * E * W + P * Y)
        t = (P * D) / denom
        t_m = t + c
        t_nom_req = t_m / (1.0 - MT)
        margin = t_act - t_m
        is_compliant = t_act >= t_m

        parameters = [
            CalculationParameter(name="Design Internal Pressure", symbol="P", value=P, unit="MPa", clause="ASME B31.3 Cl. 301.2"),
            CalculationParameter(name="Pipe Outside Diameter", symbol="D", value=D, unit="mm", clause="ASME B36.10M"),
            CalculationParameter(name="Basic Allowable Stress", symbol="S", value=S, unit="MPa", clause="ASME B31.3 Table A-1"),
            CalculationParameter(name="Casting/Weld Quality Factor", symbol="E", value=E, unit="-", clause="ASME B31.3 Table A-1A/B"),
            CalculationParameter(name="Weld Joint Strength Factor", symbol="W", value=W, unit="-", clause="ASME B31.3 Cl. 302.3.5"),
            CalculationParameter(name="Wall Thickness Coefficient", symbol="Y", value=Y, unit="-", clause="ASME B31.3 Table 304.1.1"),
            CalculationParameter(name="Mechanical Corrosion Allowance", symbol="c", value=c, unit="mm", clause="Plant Piping Spec"),
            CalculationParameter(name="Mill Fabrication Tolerance", symbol="MT", value=MT, unit="-", clause="ASTM A106 Cl. 16.2"),
            CalculationParameter(name="Measured Actual Thickness", symbol="t_act", value=t_act, unit="mm", clause="UT Inspection Survey"),
        ]

        # Excel cell mapping:
        # Inputs: C7 (P), C8 (D), C9 (S), C10 (E), C11 (W), C12 (Y), C13 (c), C14 (MT), C15 (t_act)
        # Calcs: C19 (t), C20 (t_m), C21 (t_nom_req)
        # Status: C25
        steps = [
            CalculationStep(
                name="Pressure Design Thickness",
                symbol="t",
                formula_display="t = (P * D) / (2 * (S * E * W + P * Y))",
                formula_latex=r"t = \frac{P \cdot D}{2(S \cdot E \cdot W + P \cdot Y)}",
                formula_excel="=C7*C8/(2*(C9*C10*C11+C7*C12))",
                intermediate_substitution=f"t = ({P} * {D}) / (2 * ({S} * {E} * {W} + {P} * {Y}))",
                result_value=t,
                unit="mm",
                clause="ASME B31.3 Eq. (3a)",
            ),
            CalculationStep(
                name="Minimum Required Wall Thickness",
                symbol="t_m",
                formula_display="t_m = t + c",
                formula_latex=r"t_m = t + c",
                formula_excel="=C19+C13",
                intermediate_substitution=f"t_m = {t:.4f} + {c}",
                result_value=t_m,
                unit="mm",
                clause="ASME B31.3 Cl. 304.1.1",
            ),
            CalculationStep(
                name="Minimum Required Nominal Thickness",
                symbol="t_nom_req",
                formula_display="t_nom_req = t_m / (1 - MT)",
                formula_latex=r"t_{\text{nom,req}} = \frac{t_m}{1 - MT}",
                formula_excel="=C20/(1-C14)",
                intermediate_substitution=f"t_nom_req = {t_m:.4f} / (1 - {MT})",
                result_value=t_nom_req,
                unit="mm",
                clause="ASME B31.3 Cl. 304.1.2",
            ),
        ]

        return CalculationTrace(
            title="ASME B31.3 Piping Pressure Design Thickness Verification",
            governing_standard="ASME B31.3-2022 Process Piping Clause 304.1.2",
            facility=facility,
            equipment_tag=equipment_tag,
            parameters=parameters,
            steps=steps,
            conclusion_formula='=IF(C15>=C20, "PASS - SAFE", "FAIL - UNSAFE")',
            status_condition="t_act >= t_m",
            is_compliant=is_compliant,
            safety_margin=margin,
            summary=(
                f"Evaluation of {equipment_tag} under design pressure {P} MPa. "
                f"Minimum required thickness t_m is {t_m:.3f} mm. "
                f"Actual measured thickness t_act is {t_act:.3f} mm (safety margin: {margin:+.3f} mm)."
            ),
        )
