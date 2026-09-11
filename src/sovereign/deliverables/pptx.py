"""
src/sovereign/deliverables/pptx.py
Native Microsoft PowerPoint (.pptx) executive briefing deck generator with 16:9 widescreen,
corporate styling, bulleted summaries, styled tables, native editable OOXML charts,
and Phase 1 LineageEngine packaging.
Complies with DELV-03 and T-04-04 (path fencing).
"""
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import uuid

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from sovereign.deliverables.docx import validate_output_path
from sovereign.lineage.engine import LineageEngine
from sovereign.lineage.models import (
    DeliverableInfo,
    LineageMetadata,
    ModelInfo,
)
from sovereign.tools.builtin.files import PathSecurityViolation

# Corporate Palette Constants
COLOR_PRIMARY_NAVY = RGBColor(0x1F, 0x49, 0x7D)    # #1F497D
COLOR_SECONDARY_SLATE = RGBColor(0x59, 0x59, 0x59)# #595959
COLOR_BODY_CHARCOAL = RGBColor(0x26, 0x26, 0x26)  # #262626
COLOR_LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)     # #F2F2F2
COLOR_BORDER_GRAY = RGBColor(0xD0, 0xD0, 0xD0)    # #D0D0D0
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)          # #FFFFFF
COLOR_ACCENT_BLUE = RGBColor(0x41, 0x8A, 0xB3)    # #418AB3

# 16:9 Widescreen Dimensions
SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5


class PptxDeliverableGenerator:
    """
    Generates professional Microsoft PowerPoint (.pptx) briefing decks
    featuring native editable OOXML charts, executive summary bullets,
    zebra-striped data tables, and embedded cryptographic lineage.
    """

    def __init__(self, allowed_root: Optional[Union[str, Path]] = None):
        self.allowed_root = Path(allowed_root).resolve() if allowed_root else None

    @classmethod
    def _create_presentation(cls) -> Presentation:
        """Initializes a presentation configured to 16:9 widescreen layout."""
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
        prs.slide_height = Inches(SLIDE_HEIGHT_INCHES)
        return prs

    @classmethod
    def _add_header_and_footer(
        cls,
        slide,
        heading: str,
        classification: str = "INTERNAL",
    ) -> None:
        """Adds consistent corporate top banner, heading, and bottom classification footer."""
        # Top Heading
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = heading
        p.font.name = "Calibri"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY_NAVY

        # Accent Rule (line) under header
        rule = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE
            Inches(0.8), Inches(1.25), Inches(11.733), Inches(0.04)
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = COLOR_ACCENT_BLUE
        rule.line.color.rgb = COLOR_ACCENT_BLUE

        # Bottom Footer Banner
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = f"[ {classification.upper()} ]  —  Sovereign AI Workbench  |  Confidential Industrial Briefing"
        fp.font.name = "Calibri"
        fp.font.size = Pt(9)
        fp.font.color.rgb = COLOR_SECONDARY_SLATE

    @classmethod
    def add_title_slide(
        cls,
        prs: Presentation,
        title: str,
        subtitle: str,
        classification: str = "INTERNAL",
        facility: Optional[str] = None,
        equipment_tag: Optional[str] = None,
        author: str = "Sovereign AI Workbench",
    ) -> Any:
        """Constructs an executive title slide with corporate branding and metadata."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

        # Top classification badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.4))
        btf = badge_box.text_frame
        bp = btf.paragraphs[0]
        bp.text = f"[ {classification.upper()} ]"
        bp.font.name = "Calibri"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = COLOR_SECONDARY_SLATE

        # Navy Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(2.2))
        ttf = title_box.text_frame
        ttf.word_wrap = True

        tp = ttf.paragraphs[0]
        tp.text = title
        tp.font.name = "Calibri"
        tp.font.size = Pt(32)
        tp.font.bold = True
        tp.font.color.rgb = COLOR_PRIMARY_NAVY

        sp = ttf.add_paragraph()
        sp.text = subtitle
        sp.font.name = "Calibri"
        sp.font.size = Pt(16)
        sp.font.color.rgb = COLOR_SECONDARY_SLATE
        sp.space_before = Pt(10)

        # Bottom Metadata Block
        meta_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.733), Inches(1.5))
        mtf = meta_box.text_frame

        p1 = mtf.paragraphs[0]
        fac_tag = f"Facility: {facility or 'N/A'}"
        if equipment_tag:
            fac_tag += f"    |    Equipment Tag: {equipment_tag}"
        p1.text = fac_tag
        p1.font.name = "Calibri"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_BODY_CHARCOAL

        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        p2 = mtf.add_paragraph()
        p2.text = f"Prepared By: {author}    |    Date: {now_str}"
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_SECONDARY_SLATE
        p2.space_before = Pt(4)

        return slide

    @classmethod
    def add_summary_slide(
        cls,
        prs: Presentation,
        heading: str,
        bullet_points: List[str],
        classification: str = "INTERNAL",
    ) -> Any:
        """Constructs an executive summary slide with bullet points."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        cls._add_header_and_footer(slide, heading, classification)

        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
        tf = tx_box.text_frame
        tf.word_wrap = True

        for idx, pt in enumerate(bullet_points):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•   {pt}"
            p.font.name = "Calibri"
            p.font.size = Pt(15)
            p.font.color.rgb = COLOR_BODY_CHARCOAL
            p.space_after = Pt(12)
            p.line_spacing = 1.25

        return slide

    @classmethod
    def add_table_slide(
        cls,
        prs: Presentation,
        heading: str,
        headers: List[str],
        rows: List[List[Any]],
        classification: str = "INTERNAL",
    ) -> Any:
        """Constructs a data table slide with zebra striping and navy headers."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        cls._add_header_and_footer(slide, heading, classification)

        num_rows = len(rows) + 1
        num_cols = len(headers)

        table_shape = slide.shapes.add_table(
            rows=num_rows,
            cols=num_cols,
            left=Inches(0.8),
            top=Inches(1.6),
            width=Inches(11.733),
            height=Inches(4.8),
        )
        tbl = table_shape.table

        # Format Header Row
        for col_idx, hdr_text in enumerate(headers):
            cell = tbl.cell(0, col_idx)
            cell.text = str(hdr_text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_WHITE
            p.alignment = PP_ALIGN.CENTER if col_idx == 0 else PP_ALIGN.LEFT

        # Format Data Rows
        for r_idx, row_vals in enumerate(rows, start=1):
            is_odd = (r_idx % 2 == 1)
            row_fill = COLOR_LIGHT_GRAY if is_odd else COLOR_WHITE

            for c_idx, val in enumerate(row_vals):
                if c_idx >= num_cols:
                    break
                cell = tbl.cell(r_idx, c_idx)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill

                p = cell.text_frame.paragraphs[0]
                p.font.name = "Calibri"
                p.font.size = Pt(10.5)
                p.font.color.rgb = COLOR_BODY_CHARCOAL
                p.alignment = PP_ALIGN.CENTER if c_idx == 0 else PP_ALIGN.LEFT

        return slide

    @classmethod
    def add_chart_slide(
        cls,
        prs: Presentation,
        heading: str,
        chart_type: Union[str, XL_CHART_TYPE],
        categories: List[str],
        series_data: Dict[str, List[Union[int, float]]],
        classification: str = "INTERNAL",
    ) -> Any:
        """
        Constructs a native OOXML chart slide using CategoryChartData.
        The underlying chart data remains natively editable within PowerPoint.
        """
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        cls._add_header_and_footer(slide, heading, classification)

        # Build CategoryChartData
        cdata = CategoryChartData()
        cdata.categories = categories

        for s_name, values in series_data.items():
            cdata.add_series(s_name, tuple(values))

        # Resolve chart type enum
        resolved_type = XL_CHART_TYPE.COLUMN_CLUSTERED
        if isinstance(chart_type, XL_CHART_TYPE):
            resolved_type = chart_type
        elif isinstance(chart_type, str):
            c_lower = chart_type.lower()
            if "bar" in c_lower:
                resolved_type = XL_CHART_TYPE.BAR_CLUSTERED
            elif "pie" in c_lower:
                resolved_type = XL_CHART_TYPE.PIE
            elif "line" in c_lower:
                resolved_type = XL_CHART_TYPE.LINE
            else:
                resolved_type = XL_CHART_TYPE.COLUMN_CLUSTERED

        # Add native chart shape
        chart_shape = slide.shapes.add_chart(
            resolved_type,
            Inches(1.2),
            Inches(1.6),
            Inches(10.9),
            Inches(5.0),
            cdata,
        )

        chart = chart_shape.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False

        return slide

    @classmethod
    def generate_briefing_deck(
        cls,
        output_path: Union[str, Path],
        title: str,
        facility: str,
        slides_content: List[Dict[str, Any]],
        chart_data: Optional[Dict[str, Any]] = None,
        lineage_metadata: Optional[LineageMetadata] = None,
        classification: str = "INTERNAL",
        author: str = "Sovereign Engineering AI",
        equipment_tag: Optional[str] = None,
        allowed_root: Optional[Union[str, Path]] = None,
    ) -> Path:
        """
        Generates a complete executive briefing deck in .pptx format.
        """
        resolved_path = validate_output_path(output_path, allowed_root)
        resolved_path.parent.mkdir(parents=True, exist_ok=True)

        prs = cls._create_presentation()

        # 1. Slide 1: Title Slide
        cls.add_title_slide(
            prs=prs,
            title=title,
            subtitle="Executive Technical Briefing & Condition Assessment",
            classification=classification,
            facility=facility,
            equipment_tag=equipment_tag,
            author=author,
        )

        # 2. Iterate and render content slides
        has_chart_in_content = False
        for slide_spec in slides_content:
            stype = slide_spec.get("type", "summary").lower()
            heading = slide_spec.get("heading", "Briefing Details")

            if stype in ("summary", "bullets", "text"):
                pts = slide_spec.get("bullet_points", slide_spec.get("points", []))
                cls.add_summary_slide(prs, heading, pts, classification)

            elif stype in ("table", "grid"):
                headers = slide_spec.get("headers", [])
                rows = slide_spec.get("rows", [])
                cls.add_table_slide(prs, heading, headers, rows, classification)

            elif stype in ("chart", "graph"):
                has_chart_in_content = True
                ctype = slide_spec.get("chart_type", "column_clustered")
                cats = slide_spec.get("categories", [])
                sdata = slide_spec.get("series", slide_spec.get("series_data", {}))
                cls.add_chart_slide(prs, heading, ctype, cats, sdata, classification)

        # 3. If top-level chart_data provided and not already included, add chart slide
        if chart_data and not has_chart_in_content:
            heading = chart_data.get("heading", "Component Criticality & Thickness Analysis")
            ctype = chart_data.get("chart_type", "column_clustered")
            cats = chart_data.get("categories", [])
            sdata = chart_data.get("series", chart_data.get("series_data", {}))
            cls.add_chart_slide(prs, heading, ctype, cats, sdata, classification)

        # Save presentation
        prs.save(resolved_path)

        # 4. Embed Lineage via LineageEngine
        meta = lineage_metadata
        if meta is None:
            meta = LineageMetadata(
                deliverable=DeliverableInfo(
                    file_name=resolved_path.name,
                    media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
                session_id=f"session-{uuid.uuid4().hex[:8]}",
                user_id="engineer",
                user_role="Lead Engineer",
                model=ModelInfo(model_id="sovereign-local-v1", weights_sha256="0" * 64),
                audit_start_seq=0,
                audit_end_seq=1,
                audit_tip_hash="0" * 64,
            )

        LineageEngine.package_deliverable(resolved_path, meta)
        return resolved_path
