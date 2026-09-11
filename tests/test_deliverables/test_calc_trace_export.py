"""
tests/test_deliverables/test_calc_trace_export.py
Unit tests for CalculationTraceExporter multi-format derivations (Markdown, JSON,
Word, Excel) and ToolRegistry deliverable tool integration.
"""
import json
from pathlib import Path

import docx
import openpyxl
import pytest

from sovereign.deliverables.calc_trace import (
    CalculationTrace,
    CalculationTraceExporter,
)
from sovereign.deliverables.tools import register_deliverable_tools
from sovereign.tools.registry import ToolRegistry
from sovereign.tools.schemas import UserRole


@pytest.fixture
def asme_trace():
    return CalculationTraceExporter.create_asme_b313_trace(
        pressure_mpa=2.5,
        diameter_mm=168.3,
        stress_mpa=138.0,
        quality_factor=1.0,
        weld_factor=1.0,
        y_coeff=0.4,
        corrosion_allowance_mm=3.0,
        mill_tolerance=0.125,
        actual_thickness_mm=5.4,
        facility="Jurong Island Train B",
        equipment_tag="Line-01-P-101",
    )


def test_asme_b313_mathematical_derivation(asme_trace):
    """Verify exact ASME B31.3 math formulas, substitutions, and compliance margins."""
    # Check parameters
    p_map = {p.symbol: p for p in asme_trace.parameters}
    assert p_map["P"].value == 2.5
    assert p_map["D"].value == 168.3
    assert p_map["S"].value == 138.0
    assert p_map["c"].value == 3.0

    # Check calculated steps
    s_map = {s.symbol: s for s in asme_trace.steps}

    # t = (2.5 * 168.3) / (2 * (138 * 1 * 1 + 2.5 * 0.4)) = 420.75 / 278.0 = 1.513489...
    t_step = s_map["t"]
    assert pytest.approx(t_step.result_value, 0.001) == 1.5135
    assert "t = (P * D)" in t_step.formula_display

    # t_m = t + c = 1.5135 + 3.0 = 4.5135
    tm_step = s_map["t_m"]
    assert pytest.approx(tm_step.result_value, 0.001) == 4.5135

    # t_nom_req = t_m / (1 - MT) = 4.5135 / 0.875 = 5.1583
    nom_step = s_map["t_nom_req"]
    assert pytest.approx(nom_step.result_value, 0.001) == 5.1583

    # Safety margin: 5.4 - 4.5135 = 0.8865 mm
    assert pytest.approx(asme_trace.safety_margin, 0.001) == 0.8865
    assert asme_trace.is_compliant is True


def test_calculation_trace_to_markdown(asme_trace):
    """Verify Markdown export contains tables, steps, formulas, and status badges."""
    exporter = CalculationTraceExporter(asme_trace)
    md = exporter.to_markdown()

    assert "# ASME B31.3 Piping Pressure Design Thickness Verification" in md
    assert "Jurong Island Train B" in md
    assert "Line-01-P-101" in md
    assert "| Parameter | Symbol | Value | Unit |" in md
    assert "| Design Internal Pressure | **P** | `2.5` | MPa |" in md
    assert "### Step 2.1: Pressure Design Thickness (t)" in md
    assert "**Numerical Substitution:**" in md
    assert "**PASS - SAFE**" in md
    assert "`+0.887 mm`" in md or "`+0.886 mm`" in md or "`+0.8865 mm`" in md or "0.88" in md


def test_calculation_trace_to_json(asme_trace):
    """Verify JSON round-trip serialization of CalculationTrace."""
    exporter = CalculationTraceExporter(asme_trace)
    json_str = exporter.to_json()

    data = json.loads(json_str)
    assert data["title"] == asme_trace.title
    assert len(data["parameters"]) == 9
    assert len(data["steps"]) == 3
    assert data["is_compliant"] is True

    # Validate parsing back to Pydantic model
    reconstructed = CalculationTrace.model_validate(data)
    assert reconstructed.title == asme_trace.title
    assert reconstructed.safety_margin == asme_trace.safety_margin


def test_calculation_trace_to_excel_workbook(tmp_path, asme_trace):
    """Verify exporting calculation trace to a live Excel sheet via XlsxDeliverableGenerator."""
    exporter = CalculationTraceExporter(asme_trace)
    excel_path = tmp_path / "asme_trace.xlsx"

    res_path = exporter.to_excel_workbook(excel_path)
    assert res_path == excel_path
    assert excel_path.exists()

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Engineering Calculation"]
    assert "ASME B31.3" in ws["A1"].value or "ASME B31.3" in ws["A2"].value

    # Verify input values
    assert ws["C7"].value == 2.5
    assert ws["C8"].value == 168.3

    # Verify active formulas
    assert ws["C19"].value.startswith("=")
    assert ws["C20"].value.startswith("=")


def test_calculation_trace_append_to_docx(tmp_path, asme_trace):
    """Verify appending calculation trace to a Word document."""
    doc = docx.Document()
    exporter = CalculationTraceExporter(asme_trace)
    exporter.append_to_docx(doc)

    docx_path = tmp_path / "trace_appended.docx"
    doc.save(docx_path)
    assert docx_path.exists()

    # Reopen and inspect
    doc_reloaded = docx.Document(docx_path)
    all_text = " ".join(p.text for p in doc_reloaded.paragraphs)
    assert "Engineering Calculation Trace: ASME B31.3" in all_text
    assert "Governing Standard: ASME B31.3-2022" in all_text
    assert "PASS - COMPLIANT" in all_text

    # Tables check (Parameters table)
    assert len(doc_reloaded.tables) >= 1
    tbl = doc_reloaded.tables[0]
    hdr = [c.text for c in tbl.rows[0].cells]
    assert "Parameter" in hdr[0]
    assert "Symbol" in hdr[1]


def test_tool_registry_deliverable_tools(tmp_path):
    """Verify ToolRegistry registration and execution of generate_docx, generate_xlsx, generate_pptx."""
    registry = ToolRegistry()
    register_deliverable_tools(registry, workspace_root=tmp_path)

    # 1. Check tools are registered with ENGINEER role
    doc_tool = registry.get_tool("generate_docx")
    xls_tool = registry.get_tool("generate_xlsx")
    ppt_tool = registry.get_tool("generate_pptx")

    assert doc_tool is not None
    assert doc_tool.required_role == UserRole.ENGINEER

    assert xls_tool is not None
    assert xls_tool.required_role == UserRole.ENGINEER

    assert ppt_tool is not None
    assert ppt_tool.required_role == UserRole.ENGINEER

    # 2. Test execute generate_docx via ToolRegistry
    docx_rel = "audit_report.docx"
    docx_res = registry.execute_tool(
        "generate_docx",
        {
            "output_path": docx_rel,
            "title": "Tool Registry Word Test",
            "facility": "Plant Alpha",
            "equipment_tag": "TK-101",
            "findings": [
                {
                    "item_num": 1,
                    "component": "Manway M-1",
                    "observation": "Gasket replaced",
                    "measured_value": "Nominal",
                    "severity": "Satisfactory",
                    "citation": "[Doc: doc-1234, p.1]",
                }
            ],
        },
        user_role=UserRole.ENGINEER,
    )
    assert Path(docx_res).exists()

    # 3. Test execute generate_xlsx via ToolRegistry
    xlsx_rel = "audit_calc.xlsx"
    xlsx_res = registry.execute_tool(
        "generate_xlsx",
        {
            "output_path": xlsx_rel,
            "title": "Tool Registry Excel Test",
            "standard_ref": "ASME B31.3",
            "inputs": [{"name": "P", "symbol": "P", "value": 1.0, "unit": "bar", "notes": ""}],
            "calculations": [{"name": "Step", "symbol": "S", "formula_excel": "=C7*2", "unit": "bar", "clause": ""}],
            "conclusion_formula": '=IF(C11>0, "PASS", "FAIL")',
        },
        user_role=UserRole.ENGINEER,
    )
    assert Path(xlsx_res).exists()

    # 4. Test execute generate_pptx via ToolRegistry
    pptx_rel = "audit_deck.pptx"
    pptx_res = registry.execute_tool(
        "generate_pptx",
        {
            "output_path": pptx_rel,
            "title": "Tool Registry PPTX Test",
            "facility": "Plant Beta",
            "slides_content": [
                {
                    "type": "summary",
                    "heading": "Summary",
                    "bullet_points": ["Point 1", "Point 2"],
                }
            ],
        },
        user_role=UserRole.ENGINEER,
    )
    assert Path(pptx_res).exists()

    # 5. RBAC Permission Check: Auditor role should be denied execution
    with pytest.raises(PermissionError):
        registry.execute_tool(
            "generate_docx",
            {
                "output_path": "unauthorized.docx",
                "title": "Unauthorized",
                "facility": "Plant",
                "equipment_tag": "EQ",
                "findings": [],
            },
            user_role=UserRole.AUDITOR,
        )
